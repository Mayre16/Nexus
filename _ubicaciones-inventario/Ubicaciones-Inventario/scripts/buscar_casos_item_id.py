# -*- coding: utf-8 -*-
import openpyxl
from collections import defaultdict
from sqlalchemy import create_engine, text
from config import get_config
import sys

# Configurar encoding para Windows
if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

config = get_config()
engine = create_engine(
    config.SQLALCHEMY_DATABASE_URI,
    **config.SQLALCHEMY_ENGINE_OPTIONS
)

# Cargar el archivo Excel
archivo = 'EJEMPLO_AJUSTES_MASIVOS (1).xlsx'
wb = openpyxl.load_workbook(archivo)
ws = wb.active

# Leer SKUs y ubicaciones del Excel
filas_excel = []
for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
    if any(cell.value for cell in row):
        sku = str(row[0].value).strip().upper() if row[0].value else ''
        ubicacion = str(row[1].value).strip().upper() if row[1].value else ''
        cantidad = row[2].value if len(row) > 2 else None
        product_id_excel = str(row[4].value).strip() if len(row) > 4 and row[4].value else ''
        
        if sku and ubicacion:
            try:
                cantidad_float = float(cantidad) if cantidad else 0.0
                if cantidad_float > 0:
                    filas_excel.append({
                        'fila': row_idx,
                        'sku': sku,
                        'ubicacion': ubicacion,
                        'cantidad': cantidad_float,
                        'product_id_excel': product_id_excel
                    })
            except:
                pass

print(f"=== ANALISIS DEL ARCHIVO EXCEL ===")
print(f"Total filas validas: {len(filas_excel)}")
print()

# Consultar la BD para obtener item_ids de estos SKUs
print("Consultando base de datos para obtener item_ids...")
skus_unicos = list(set([f['sku'] for f in filas_excel]))

with engine.connect() as conn:
    # Buscar item_ids para cada SKU
    sku_a_item_id = {}
    sku_a_nombre = {}
    
    # Consultar en lotes para evitar queries muy largas
    batch_size = 100
    for i in range(0, len(skus_unicos), batch_size):
        batch = skus_unicos[i:i+batch_size]
        # Crear lista de placeholders
        placeholders = ','.join([f':sku_{j}' for j in range(len(batch))])
        # Crear diccionario de parámetros
        params = {f'sku_{j}': sku for j, sku in enumerate(batch)}
        
        query = text(f"""
            SELECT sku, item_id, nombre 
            FROM productos_adm 
            WHERE sku IN ({placeholders})
        """)
        
        result = conn.execute(query, params)
        for row in result:
            sku_a_item_id[row[0]] = row[1]
            sku_a_nombre[row[0]] = row[2] or 'Sin nombre'
    
    print(f"SKUs encontrados en BD: {len(sku_a_item_id)} de {len(skus_unicos)}")
    print()
    
    # Agregar item_id a cada fila del Excel
    for fila in filas_excel:
        sku = fila['sku']
        if sku in sku_a_item_id:
            fila['item_id'] = sku_a_item_id[sku]
            fila['nombre'] = sku_a_nombre[sku]
        else:
            # Si no se encuentra, usar el product_id del Excel o el SKU como fallback
            fila['item_id'] = fila['product_id_excel'] if fila['product_id_excel'] else sku
            fila['nombre'] = 'No encontrado en BD'
    
    # Agrupar por (item_id, ubicacion) - igual que hace el código
    agrupados = defaultdict(list)
    for fila in filas_excel:
        key = (fila['item_id'], fila['ubicacion'])
        agrupados[key].append(fila)
    
    print(f"=== RESULTADO DEL AGRUPAMIENTO ===")
    print(f"Filas en Excel: {len(filas_excel)}")
    print(f"Combinaciones unicas (item_id, ubicacion): {len(agrupados)}")
    print(f"Diferencia: {len(filas_excel) - len(agrupados)} filas se agruparon")
    print()
    
    # Buscar casos donde múltiples filas se agruparon
    casos_agrupados = {k: v for k, v in agrupados.items() if len(v) > 1}
    
    if casos_agrupados:
        print(f"=== CASOS DONDE MULTIPLES FILAS SE AGRUPARON ===")
        print(f"Total casos: {len(casos_agrupados)}")
        print()
        
        # Ordenar por cantidad de filas agrupadas
        casos_ordenados = sorted(casos_agrupados.items(), key=lambda x: len(x[1]), reverse=True)
        
        # Mostrar los primeros 5 casos
        for i, ((item_id, ubicacion), filas) in enumerate(casos_ordenados[:5], 1):
            print(f"--- CASO {i} ---")
            print(f"item_id: {item_id}")
            print(f"Ubicacion: {ubicacion}")
            print(f"Filas agrupadas: {len(filas)}")
            
            # Obtener SKUs únicos en este grupo
            skus_en_grupo = list(set([f['sku'] for f in filas]))
            print(f"SKUs diferentes: {len(skus_en_grupo)}")
            print(f"Lista de SKUs: {', '.join(skus_en_grupo[:10])}")
            if len(skus_en_grupo) > 10:
                print(f"  ... y {len(skus_en_grupo) - 10} SKUs mas")
            
            # Mostrar nombres de productos
            print(f"Productos:")
            for sku in skus_en_grupo[:5]:
                nombre = filas[0]['nombre'] if filas else 'N/A'
                # Buscar el nombre correcto para este SKU
                for f in filas:
                    if f['sku'] == sku:
                        nombre = f['nombre']
                        break
                print(f"  - {sku}: {nombre[:60]}")
            
            # Mostrar algunas filas
            print(f"Ejemplo de filas (primeras 5):")
            for fila in filas[:5]:
                print(f"  Fila {fila['fila']}: SKU={fila['sku']}, Cantidad={fila['cantidad']}")
            if len(filas) > 5:
                print(f"  ... y {len(filas) - 5} filas mas")
            print()
    else:
        print("No se encontraron casos de agrupamiento.")
        print("Esto significa que cada (item_id, ubicacion) es unico.")
        print()
        print("La diferencia entre 4500 y 3656 debe deberse a:")
        print("1. Filas con SKUs no encontrados en la BD")
        print("2. Filas con ubicaciones fisicas que no existen o estan inactivas")
        print("3. Filas con cantidad <= 0")
